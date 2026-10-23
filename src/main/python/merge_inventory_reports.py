import os
import sys
import json
import re
import argparse
from collections import Counter
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Fix Windows terminal UTF-8 encoding
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

def merge_inventory_reports(file1_path: str, file2_path: str, output_path: str = None) -> dict:
    """
    Merges two Excel Inventory Reports into a single combined report workbook.
    """
    if not os.path.exists(file1_path):
        raise FileNotFoundError(f"File 1 not found: {file1_path}")
    if not os.path.exists(file2_path):
        raise FileNotFoundError(f"File 2 not found: {file2_path}")

    if not output_path:
        dir_name = os.path.dirname(file1_path)
        base1 = os.path.basename(file1_path)
        if '_01M' in base1:
            out_name = base1.replace('_01M', '_Merged')
        elif '_02E' in base1:
            out_name = base1.replace('_02E', '_Merged')
        else:
            out_name = os.path.splitext(base1)[0] + '_Merged.xlsx'
        output_path = os.path.join(dir_name, out_name)

    wb1 = openpyxl.load_workbook(file1_path, data_only=True)
    wb2 = openpyxl.load_workbook(file2_path, data_only=True)

    out_wb = openpyxl.Workbook()
    # Remove default initial sheet
    out_wb.remove(out_wb.active)

    # Styling definitions
    header_fill_fs = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
    header_font_fs = Font(name='Calibri', size=11, bold=True, color='001F4E78')

    sub_fill_fs = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    sub_font_fs = Font(name='Calibri', size=10, bold=True)

    missing_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    missing_font = Font(name='Calibri', size=11, bold=True, color='9C0006')

    bold_font = Font(name='Calibri', size=11, bold=True)

    stats = {
        'missingCount': 0,
        'folderColsCount': 0,
        'cameraSummary': {},
        'issueCount': 0
    }

    # 1. Missing Files Sheet
    ws_mf = out_wb.create_sheet(title='Missing Files')
    mf_headers = ['Folder Path', 'File Name', 'Status', 'Camera Type']
    ws_mf.append(mf_headers)
    for c in range(1, 5):
        ws_mf.cell(1, c).font = bold_font

    mf_rows = []
    for wb in [wb1, wb2]:
        if 'Missing Files' in wb.sheetnames:
            ws = wb['Missing Files']
            for r in range(2, ws.max_row + 1):
                row_vals = [ws.cell(r, c).value for c in range(1, 5)]
                if any(v is not None for v in row_vals):
                    mf_rows.append(row_vals)

    for r_idx, row in enumerate(mf_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_mf.cell(r_idx, c_idx, val)
            cell.fill = missing_fill
            cell.font = missing_font
    stats['missingCount'] = len(mf_rows)

    # 2. Folder Structure Sheet
    ws_fs = out_wb.create_sheet(title='Folder Structure')
    def extract_fs_cols(wb):
        if 'Folder Structure' not in wb.sheetnames: return []
        ws = wb['Folder Structure']
        cols = []
        for c in range(1, ws.max_column + 1):
            r1 = ws.cell(1, c).value
            r2 = ws.cell(2, c).value
            if r1 is None and r2 is None: continue
            vals = [ws.cell(r, c).value for r in range(3, ws.max_row + 1)]
            files = []
            for v in vals:
                if v is None: continue
                sv = str(v).strip()
                if not sv: continue
                if re.match(r'^\.[a-zA-Z0-9]+:\s*\d+$', sv): continue # Skip old summary lines
                files.append(sv)
            cols.append((r1, r2, files))
        return cols

    cols1 = extract_fs_cols(wb1)
    cols2 = extract_fs_cols(wb2)

    fs_merged = {}
    fs_order = []

    for r1, r2, files in cols1:
        key = (r1, r2)
        if key not in fs_merged:
            fs_merged[key] = []
            fs_order.append(key)
        fs_merged[key].extend(files)

    for r1, r2, files in cols2:
        key = (r1, r2)
        if key not in fs_merged:
            fs_merged[key] = []
            fs_order.append(key)
        fs_merged[key].extend(files)

    for col_idx, (r1, r2) in enumerate(fs_order, start=1):
        c1 = ws_fs.cell(1, col_idx, r1)
        c1.fill = header_fill_fs
        c1.font = header_font_fs
        
        c2 = ws_fs.cell(2, col_idx, r2)
        c2.fill = sub_fill_fs
        c2.font = sub_font_fs
        
        files = fs_merged[(r1, r2)]
        row_curr = 3
        ext_counter = Counter()
        for f in files:
            ws_fs.cell(row_curr, col_idx, f)
            row_curr += 1
            m = re.search(r'\.([a-zA-Z0-9]+)$', f)
            if m:
                ext_counter[m.group(1).lower()] += 1
                
        for ext, count in sorted(ext_counter.items()):
            ws_fs.cell(row_curr, col_idx, f'.{ext}: {count}')
            row_curr += 1

    stats['folderColsCount'] = len(fs_order)

    # 3. Camera Summary Sheet
    ws_cs = out_wb.create_sheet(title='Camera Summary')
    def extract_cs_cols(wb):
        if 'Camera Summary' not in wb.sheetnames: return {}
        ws = wb['Camera Summary']
        result = {}
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        for c_idx, h in enumerate(headers, start=1):
            if not h: continue
            vals = [ws.cell(r, c_idx).value for r in range(2, ws.max_row + 1)]
            files = []
            tot_found = None
            tot_missing = None
            for v in vals:
                if v is None: continue
                sv = str(v).strip()
                if sv.startswith('Total Found:'):
                    try: tot_found = int(sv.split(':')[1].strip())
                    except: pass
                elif sv.startswith('Total Missing:'):
                    try: tot_missing = int(sv.split(':')[1].strip())
                    except: pass
                else:
                    files.append(sv)
            result[h] = {'files': files, 'tot_found': tot_found, 'tot_missing': tot_missing}
        return result

    cs1 = extract_cs_cols(wb1)
    cs2 = extract_cs_cols(wb2)

    all_cs_headers = list(cs1.keys())
    for h in cs2.keys():
        if h not in all_cs_headers:
            all_cs_headers.append(h)

    for col_idx, h in enumerate(all_cs_headers, start=1):
        cell_h = ws_cs.cell(1, col_idx, h)
        cell_h.font = bold_font
        
        files1 = cs1.get(h, {}).get('files', [])
        files2 = cs2.get(h, {}).get('files', [])
        tf1 = cs1.get(h, {}).get('tot_found') or len(files1)
        tf2 = cs2.get(h, {}).get('tot_found') or len(files2)
        tm1 = cs1.get(h, {}).get('tot_missing') or 0
        tm2 = cs2.get(h, {}).get('tot_missing') or 0
        
        combined_files = files1 + files2
        row_curr = 2
        for f in combined_files:
            ws_cs.cell(row_curr, col_idx, f)
            row_curr += 1
            
        found_sum = tf1 + tf2
        missing_sum = tm1 + tm2
        
        c_tf = ws_cs.cell(row_curr, col_idx, f'Total Found: {found_sum}')
        c_tf.font = bold_font
        row_curr += 1
        c_tm = ws_cs.cell(row_curr, col_idx, f'Total Missing: {missing_sum}')
        c_tm.font = bold_font

        stats['cameraSummary'][h] = {'found': found_sum, 'missing': missing_sum}

    # 4. Contributor Sheet
    ws_con = out_wb.create_sheet(title='Contributor')
    ws_con.append(['Folder Path', 'File Name'])
    for c in range(1, 3): ws_con.cell(1, c).font = bold_font
    con_rows = []
    for wb in [wb1, wb2]:
        if 'Contributor' in wb.sheetnames:
            ws = wb['Contributor']
            for r in range(2, ws.max_row + 1):
                rv = [ws.cell(r, c).value for c in range(1, 3)]
                if any(v is not None for v in rv) and rv not in con_rows:
                    con_rows.append(rv)
    for r_idx, r in enumerate(con_rows, start=2):
        ws_con.cell(r_idx, 1, r[0])
        ws_con.cell(r_idx, 2, r[1])

    # 5. Issues Sheet
    ws_iss = out_wb.create_sheet(title='Issues')
    ws_iss.append(['Path', 'Issue Type', 'Description'])
    for c in range(1, 4): ws_iss.cell(1, c).font = bold_font
    iss_rows = []
    for wb in [wb1, wb2]:
        if 'Issues' in wb.sheetnames:
            ws = wb['Issues']
            for r in range(2, ws.max_row + 1):
                rv = [ws.cell(r, c).value for c in range(1, 4)]
                if any(v is not None for v in rv) and rv not in iss_rows:
                    iss_rows.append(rv)
    for r_idx, r in enumerate(iss_rows, start=2):
        for c_idx, val in enumerate(r, start=1):
            ws_iss.cell(r_idx, c_idx, val)
    stats['issueCount'] = len(iss_rows)

    # Auto-adjust column widths
    for sheet in out_wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 100)

    out_wb.save(output_path)
    return {
        'success': True,
        'outputPath': output_path,
        'stats': stats
    }

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Merge two Excel Inventory Reports into a single merged workbook.')
    parser.add_argument('--file1', required=True, help='Path to Report 1 (.xlsx)')
    parser.add_argument('--file2', required=True, help='Path to Report 2 (.xlsx)')
    parser.add_argument('--output', '-o', help='Path for output merged (.xlsx)')
    parser.add_argument('--json', action='store_true', help='Output status as JSON')

    args = parser.parse_args()
    try:
        res = merge_inventory_reports(args.file1, args.file2, args.output)
        if args.json:
            print(json.dumps(res))
        else:
            print(f"Successfully merged inventory reports into: {res['outputPath']}")
    except Exception as e:
        err_res = {'success': False, 'error': str(e)}
        if args.json:
            print(json.dumps(err_res))
        else:
            print(f"Error merging inventory reports: {e}", file=sys.stderr)
        sys.exit(1)
