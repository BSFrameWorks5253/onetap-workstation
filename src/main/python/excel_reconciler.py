#!/usr/bin/env python3
"""
============================================================
EXCEL RECONCILIATION ENGINE (src/main/python/excel_reconciler.py)
Part of One Tap Workstation

Audits current production spreadsheets (.xlsx) against historical
scanned archives. Resolves date folder discrepancies using dual-calendar
(Hijri/Gregorian) canonical normalization and multi-tier heuristics.
============================================================
"""

import sys
import os
import re
import json
import sqlite3
import argparse
from datetime import datetime, date
from typing import Dict, List, Any, Optional, Tuple

# Fix Windows terminal UTF-8 encoding
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

script_dir = os.path.dirname(os.path.abspath(__file__))
if script_dir not in sys.path:
    sys.path.insert(0, script_dir)

import archive_catalog_db as db_mod
import archive_scanner as scanner_mod

# -------------------------------------------------------------------
# Dual-Calendar & Session Normalizers
# -------------------------------------------------------------------

HIJRI_MONTHS = [
    'moharram', 'safar', 'rabi i', 'rabi ii', 'jumada i', 'jumada ii',
    'rajab', 'shaban', 'ramadan', 'shawwal', 'zilqad', 'zilhajj',
    "rabi' al-awwal", "rabi' al-aakhar", "jumada al-ula", "jumada al-ukhra"
]

def normalize_session(text: str) -> Optional[str]:
    """Maps varied session descriptions to canonical session tags."""
    if not text:
        return None
    s = str(text).strip().upper()
    if '_01M' in s or 'MORNING' in s or 'SUBAH' in s or 'S1' in s:
        return '_01M'
    if '_02E' in s or 'EVENING' in s or 'RAAT' in s or 'NIGHT' in s or 'S2' in s:
        return '_02E'
    return None

def normalize_date_input(val: Any) -> Dict[str, Any]:
    """
    Parses and canonicalizes varied date representations into YYYY-MM-DD.
    Handles datetime objects, Hijri strings, and standard date strings.
    """
    res = {
        'hijri_canonical': None,
        'greg_canonical': None,
        'raw': str(val).strip() if val is not None else ''
    }
    if not val:
        return res

    # 1. Direct Python datetime/date object from openpyxl
    if isinstance(val, (datetime, date)):
        res['greg_canonical'] = val.strftime('%Y-%m-%d')
        return res

    s = str(val).strip()

    # 2. Check for Hijri YYYY-MM-DD pattern (years 1400-1500)
    m_hijri = re.search(r'(14\d{2})[\s._\-](\d{1,2})[\s._\-](\d{1,2})', s)
    if m_hijri:
        y, m, d = int(m_hijri.group(1)), int(m_hijri.group(2)), int(m_hijri.group(3))
        if 1 <= m <= 12 and 1 <= d <= 31:
            res['hijri_canonical'] = f"{y:04d}-{m:02d}-{d:02d}"
            return res

    # 3. Check for Gregorian YYYY-MM-DD pattern (years 2000-2099)
    m_greg = re.search(r'(20\d{2})[\s._\-](\d{1,2})[\s._\-](\d{1,2})', s)
    if m_greg:
        y, m, d = int(m_greg.group(1)), int(m_greg.group(2)), int(m_greg.group(3))
        if 1 <= m <= 12 and 1 <= d <= 31:
            res['greg_canonical'] = f"{y:04d}-{m:02d}-{d:02d}"
            return res

    # 4. DD/MM/YYYY or DD-MM-YYYY format
    m_dmy = re.search(r'(\d{1,2})[/\-.](\d{1,2})[/\-.](20\d{2})', s)
    if m_dmy:
        d, m, y = int(m_dmy.group(1)), int(m_dmy.group(2)), int(m_dmy.group(3))
        if 1 <= m <= 12 and 1 <= d <= 31:
            res['greg_canonical'] = f"{y:04d}-{m:02d}-{d:02d}"
            return res

    # 5. Month name parsing (e.g. "13 Rabi I 1448")
    s_lower = s.lower()
    for idx, month_name in enumerate(HIJRI_MONTHS, start=1):
        if month_name in s_lower:
            month_num = (idx if idx <= 12 else (idx - 12))
            dm = re.search(r'(\d{1,2})', s)
            ym = re.search(r'(14\d{2})', s)
            d = int(dm.group(1)) if dm else 1
            y = int(ym.group(1)) if ym else 1448
            res['hijri_canonical'] = f"{y:04d}-{month_num:02d}-{d:02d}"
            return res

    return res

# -------------------------------------------------------------------
# Core Reconciliation Engine
# -------------------------------------------------------------------

def reconcile_excel_against_archive(
    excel_path: str,
    archive_id: int,
    sheet_name: Optional[str] = None,
    col_date: Optional[str] = None,
    col_folder: Optional[str] = None,
    col_file: Optional[str] = None,
    date_filter: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    db_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Compares an Excel file against a scanned archive in SQLite.
    Returns categorized match records, folder-by-folder count breakdowns, and summary KPIs.
    """
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    conn = db_mod.get_connection(db_path)
    cur = conn.cursor()

    # Verify archive exists
    cur.execute("SELECT * FROM scanned_archives WHERE id = ?", (archive_id,))
    archive = cur.fetchone()
    if not archive:
        conn.close()
        raise ValueError(f"Archive ID {archive_id} not found in database.")

    # Query all media items in this archive with optional date filter
    query = """
        SELECT 
            m.id as media_id,
            m.file_name,
            m.file_stem,
            m.size_bytes,
            m.duration_seconds,
            m.duration_formatted,
            f.id as folder_id,
            f.relative_path,
            f.folder_name,
            f.session_tag,
            f.session_name,
            f.event_code,
            f.date_canonical_hijri,
            f.date_canonical_greg
        FROM media_items m
        JOIN folder_nodes f ON f.id = m.folder_id
        WHERE m.archive_id = ?
    """
    params: List[Any] = [archive_id]
    if date_filter and date_filter.strip():
        df = date_filter.strip()
        query += " AND (f.date_canonical_hijri = ? OR f.date_canonical_greg = ? OR f.relative_path LIKE ?)"
        params.extend([df, df, f"%{df}%"])
    elif date_from and date_to:
        query += " AND ((f.date_canonical_hijri BETWEEN ? AND ?) OR (f.date_canonical_greg BETWEEN ? AND ?))"
        params.extend([date_from, date_to, date_from, date_to])

    cur.execute(query, params)
    media_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    # Index media items by filename and by stem
    by_filename: Dict[str, List[Dict[str, Any]]] = {}
    by_stem: Dict[str, List[Dict[str, Any]]] = {}
    archive_folders: Dict[str, Dict[str, Any]] = {}

    for item in media_rows:
        fn_key = item['file_name'].lower()
        st_key = item['file_stem'].lower()
        by_filename.setdefault(fn_key, []).append(item)
        by_stem.setdefault(st_key, []).append(item)

        rp = item['relative_path']
        if rp not in archive_folders:
            archive_folders[rp] = {
                'relative_path': rp,
                'files': set(),
                'mp4_count': 0
            }
        archive_folders[rp]['files'].add(fn_key)
        if fn_key.endswith(('.mp4', '.mov')):
            archive_folders[rp]['mp4_count'] += 1

    # Load Workbook
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    # Detect Header Columns if not explicitly passed
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    date_col_idx = 1
    folder_col_idx = 2
    file_col_idx = 3

    for idx, h in enumerate(headers, start=1):
        if h:
            hs = str(h).strip().lower()
            if col_date and col_date.lower() in hs:
                date_col_idx = idx
            elif col_folder and col_folder.lower() in hs:
                folder_col_idx = idx
            elif col_file and col_file.lower() in hs:
                file_col_idx = idx
            else:
                if 'date' in hs or 'tarikh' in hs:
                    date_col_idx = idx
                elif 'folder' in hs or 'session' in hs or 'category' in hs or 'event' in hs:
                    folder_col_idx = idx
                elif 'file' in hs or 'clip' in hs or 'name' in hs or 'video' in hs:
                    file_col_idx = idx

    results = []
    matched_exact = 0
    matched_shifted = 0
    missing_count = 0
    excel_folders: Dict[str, Dict[str, Any]] = {}

    # Iterate rows starting from row 2
    for r in range(2, ws.max_row + 1):
        raw_date = ws.cell(r, date_col_idx).value
        raw_folder = ws.cell(r, folder_col_idx).value
        raw_file = ws.cell(r, file_col_idx).value

        # Skip blank lines
        if raw_date is None and raw_folder is None and raw_file is None:
            continue

        file_str = str(raw_file).strip() if raw_file is not None else ''
        if not file_str:
            continue

        folder_str = str(raw_folder).strip() if raw_folder is not None else 'Root'
        if folder_str not in excel_folders:
            excel_folders[folder_str] = {
                'folder_name': folder_str,
                'raw_date': str(raw_date).strip() if raw_date is not None else '',
                'excel_count': 0,
                'matched_count': 0,
                'missing_count': 0,
                'matched_archive_paths': set()
            }
        excel_folders[folder_str]['excel_count'] += 1

        # Clean filename / stem
        file_stem, _ = os.path.splitext(file_str)
        date_norm = normalize_date_input(raw_date)
        session_norm = normalize_session(raw_folder)

        # 1. Candidate lookup by exact filename or stem
        candidates = by_filename.get(file_str.lower(), [])
        if not candidates:
            candidates = by_stem.get(file_stem.lower(), [])

        record = {
            'row_index': r,
            'raw_date': str(raw_date) if raw_date is not None else '',
            'raw_folder': folder_str,
            'raw_file': file_str,
            'match_status': 'MISSING',
            'confidence': 0.0,
            'explanation': 'File not found anywhere in this Safar archive.',
            'matched_item': None
        }

        if candidates:
            best_match = None
            best_score = 0.0
            best_status = 'FOLDER_SHIFT'
            best_explanation = 'Found in archive, but date or session folder differs.'

            for cand in candidates:
                cand_hijri = cand.get('date_canonical_hijri')
                cand_greg = cand.get('date_canonical_greg')
                cand_session = cand.get('session_tag')

                # Check exact date match
                date_matches = (
                    (date_norm['hijri_canonical'] and cand_hijri == date_norm['hijri_canonical']) or
                    (date_norm['greg_canonical'] and cand_greg == date_norm['greg_canonical'])
                )

                # Check session match
                session_matches = (session_norm and cand_session == session_norm)

                if date_matches and session_matches:
                    best_match = cand
                    best_score = 1.0
                    best_status = 'EXACT'
                    best_explanation = 'Exact match: date, session, and file name align.'
                    break
                elif date_matches:
                    if best_score < 0.85:
                        best_match = cand
                        best_score = 0.85
                        best_status = 'SESSION_SHIFT'
                        best_explanation = f"Date matches, but session in archive is '{cand_session or 'N/A'}'."
                elif best_score < 0.65:
                    best_match = cand
                    best_score = 0.65
                    best_status = 'FOLDER_SHIFT'
                    best_explanation = f"Found in folder '{cand['relative_path']}' with date '{cand_hijri or cand_greg or 'N/A'}'."

            record['match_status'] = best_status
            record['confidence'] = best_score
            record['explanation'] = best_explanation
            record['matched_item'] = {
                'media_id': best_match['media_id'],
                'file_name': best_match['file_name'],
                'folder_path': best_match['relative_path'],
                'size_bytes': best_match['size_bytes'],
                'duration_formatted': best_match['duration_formatted']
            }

            excel_folders[folder_str]['matched_count'] += 1
            excel_folders[folder_str]['matched_archive_paths'].add(best_match['relative_path'])

            if best_status == 'EXACT':
                matched_exact += 1
            else:
                matched_shifted += 1
        else:
            missing_count += 1
            excel_folders[folder_str]['missing_count'] += 1

        results.append(record)

    # -----------------------------------------------------------------
    # Folder-by-Folder Breakdown & Counts
    # -----------------------------------------------------------------
    folder_summary = []
    accounted_archive_folders = set()

    for f_name, f_data in excel_folders.items():
        # Match archive folder
        arch_match = archive_folders.get(f_name)
        if not arch_match and f_data['matched_archive_paths']:
            best_rel = next(iter(f_data['matched_archive_paths']))
            arch_match = archive_folders.get(best_rel)
        if not arch_match:
            for k, v in archive_folders.items():
                if k.lower() == f_name.lower() or k.lower().endswith(f_name.lower()) or f_name.lower().endswith(k.lower()):
                    arch_match = v
                    break

        arch_file_count = len(arch_match['files']) if arch_match else 0
        if arch_match:
            accounted_archive_folders.add(arch_match['relative_path'])

        exc_count = f_data['excel_count']
        diff = exc_count - arch_file_count

        if arch_file_count == 0:
            status = 'MISSING_IN_ARCHIVE'
        elif diff == 0 and f_data['missing_count'] == 0:
            status = 'BALANCED'
        else:
            status = 'COUNT_MISMATCH'

        folder_summary.append({
            'folder_name': f_name,
            'matched_archive_path': arch_match['relative_path'] if arch_match else '-',
            'date_val': f_data['raw_date'],
            'excel_file_count': exc_count,
            'archive_file_count': arch_file_count,
            'matched_files': f_data['matched_count'],
            'missing_files': f_data['missing_count'],
            'difference': diff,
            'status': status
        })

    # Archive folders with 0 files in Excel
    for arch_path, arch_info in archive_folders.items():
        if arch_path not in accounted_archive_folders:
            folder_summary.append({
                'folder_name': arch_path,
                'matched_archive_path': arch_path,
                'date_val': '',
                'excel_file_count': 0,
                'archive_file_count': len(arch_info['files']),
                'matched_files': 0,
                'missing_files': 0,
                'difference': -len(arch_info['files']),
                'status': 'EXTRA_IN_ARCHIVE'
            })

    # Sort folder breakdown: Mismatches/Missing first, then Balanced
    status_order = {'MISSING_IN_ARCHIVE': 0, 'COUNT_MISMATCH': 1, 'EXTRA_IN_ARCHIVE': 2, 'BALANCED': 3}
    folder_summary.sort(key=lambda f: (status_order.get(f['status'], 9), f['folder_name']))

    balanced_f_count = sum(1 for f in folder_summary if f['status'] == 'BALANCED')
    mismatch_f_count = sum(1 for f in folder_summary if f['status'] == 'COUNT_MISMATCH')
    missing_f_count = sum(1 for f in folder_summary if f['status'] == 'MISSING_IN_ARCHIVE')
    extra_f_count = sum(1 for f in folder_summary if f['status'] == 'EXTRA_IN_ARCHIVE')

    summary = {
        'total_excel_rows': len(results),
        'total_archive_files': len(media_rows),
        'total_excel_folders': len(excel_folders),
        'total_archive_folders': len(archive_folders),
        'balanced_folder_count': balanced_f_count,
        'mismatched_folder_count': mismatch_f_count,
        'missing_folder_count': missing_f_count,
        'extra_folder_count': extra_f_count,
        'matched_exact_count': matched_exact,
        'matched_shifted_count': matched_shifted,
        'missing_count': missing_count,
        'accuracy_rate': round((matched_exact + matched_shifted) / max(len(results), 1) * 100, 1)
    }

    return {
        'success': True,
        'safar_name': archive['safar_name'],
        'summary': summary,
        'folder_summary': folder_summary,
        'records': results
    }

# -------------------------------------------------------------------
# Export Formatted Audit Report (.xlsx)
# -------------------------------------------------------------------

def export_audit_excel(audit_data: Dict[str, Any], output_path: str) -> str:
    """Generates an executive color-coded reconciliation spreadsheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Reconciliation"

    # Styling fills
    exact_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")     # Soft Green
    shifted_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")   # Soft Yellow
    missing_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")   # Soft Red
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")    # Dark Slate

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)

    headers = [
        "Row", "Excel Date", "Excel Folder / Session", "Excel Filename",
        "Audit Status", "Confidence", "Archive Folder Path", "Archive File", "Duration", "Explanation"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for rec in audit_data.get('records', []):
        m = rec.get('matched_item') or {}
        row_vals = [
            rec['row_index'],
            rec['raw_date'],
            rec['raw_folder'],
            rec['raw_file'],
            rec['match_status'],
            f"{int(rec['confidence'] * 100)}%",
            m.get('folder_path', '-'),
            m.get('file_name', '-'),
            m.get('duration_formatted', '-'),
            rec['explanation']
        ]
        ws.append(row_vals)
        cur_row = ws.max_row

        status = rec['match_status']
        row_fill = exact_fill if status == 'EXACT' else (shifted_fill if 'SHIFT' in status else missing_fill)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(cur_row, col_idx)
            c.fill = row_fill
            c.font = regular_font

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    wb.save(output_path)
    return output_path

OFFICIAL_EVENTS_LIST = [
    ('AQ', 'Aqiqa'),
    ('BM', 'Bethak Majlis'),
    ('BQ', 'Bethak Qadam'),
    ('BZ', 'Bethak Ziyafat'),
    ('FMB', 'FMB'),
    ('H', 'Haddiyat'),
    ('ID', 'Idaraat'),
    ('IF', 'Iftettah'),
    ('IFS', 'Iftettah Sharaf'),
    ('IST', 'Isteqbal'),
    ('KGS', 'Kg Sharaf'),
    ('M', 'Majlis'),
    ('MF', 'Mafsuhiyat'),
    ('MISBAH', 'Misbah'),
    ('MKB', 'Mokib'),
    ('MO', 'Madrasa'),
    ('MQ', 'Misaq'),
    ('MUL', 'Mulaqaat'),
    ('MYN', 'Miyana'),
    ('MZ', 'Majmui Ziyafat'),
    ('N', 'Nikah'),
    ('NAMAZ', 'Namaz'),
    ('ONLINE', 'Online'),
    ('Q', 'Qadam'),
    ('QMZ', 'Qadam Majlis Ziyafat'),
    ('SH', 'Sharaf'),
    ('SQ', 'Shifa Qadam'),
    ('TJR', 'Tijarat'),
    ('TQ', 'Talaqqi'),
    ('TQN', 'Tahneeq'),
    ('WAAZ', 'Waaz'),
    ('WD', 'Wada'),
    ('Z', 'Ziyafat'),
    ('ZO', 'Zafaf Olda')
]

OFFICIAL_EVENTS_DICT = dict(OFFICIAL_EVENTS_LIST)

def parse_duration_to_seconds(dur_str: str) -> float:
    if not dur_str:
        return 0.0
    parts = str(dur_str).strip().split(':')
    try:
        if len(parts) == 3:
            return float(parts[0]) * 3600 + float(parts[1]) * 60 + float(parts[2])
        elif len(parts) == 2:
            return float(parts[0]) * 60 + float(parts[1])
        elif len(parts) == 1:
            return float(parts[0])
    except Exception:
        return 0.0
    return 0.0

def format_seconds_to_duration(seconds: float) -> str:
    if not seconds or seconds <= 0:
        return "00:00"
    total_sec = int(round(seconds))
    hrs = total_sec // 3600
    rem = total_sec % 3600
    mins = rem // 60
    secs = rem % 60
    if hrs > 0:
        return f"{hrs:02d}:{mins:02d}:{secs:02d}"
    return f"{mins:02d}:{secs:02d}"

def sanitize_sheet_title(name: str) -> str:
    cleaned = re.sub(r'[\\/\*\[\]\:\?]', '_', name).strip()
    return cleaned[:31] if len(cleaned) > 31 else (cleaned or "Sheet")

def build_segregated_workbook(
    title: str,
    subtitle: str,
    items: List[Dict[str, Any]],
    output_path: str
) -> Dict[str, Any]:
    """
    Builds a studio-grade multi-tab segregated Excel workbook:
      1. '📊 Summary & KPIs' (KPI cards, Event matrix, Date breakdown)
      2. Event-specific segregated sheets (e.g. 'Q - Qadam', 'M - Majlis', 'N - Nikah'...)
      3. 'Full Audit Sheet' (Complete master list for reconciler/filters)
      4. 'Official Event Codes' (Legend)
    """
    wb = openpyxl.Workbook()

    # Reusable Styles
    brand_navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    brand_blue_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    dark_slate_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    kpi_card_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    section_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_banner = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_subbanner = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    font_section = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_kpi_label = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    font_kpi_value = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    data_font = Font(name="Segoe UI", size=10)
    total_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    total_border = Border(
        top=Side(style='thin', color='94A3B8'),
        bottom=Side(style='double', color='0F172A'),
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0')
    )

    fill_uploaded = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    font_uploaded = Font(name="Segoe UI", size=9, bold=True, color="065F46")
    fill_submitted = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    font_submitted = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
    fill_pending = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_pending = Font(name="Segoe UI", size=9, bold=True, color="92400E")
    fill_not_uploaded = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_not_uploaded = Font(name="Segoe UI", size=9, color="64748B")

    # Global KPI Calculations
    total_clips = len(items)
    total_seconds = sum(parse_duration_to_seconds(it.get('duration', '')) for it in items)
    total_duration_fmt = format_seconds_to_duration(total_seconds)
    total_size_mb = sum((it.get('size_mb') or 0.0) for it in items)
    total_size_gb = round(total_size_mb / 1024, 2)

    status_counts = {'UPLOADED': 0, 'SUBMITTED': 0, 'PENDING UPLOAD': 0, 'NOT UPLOADED': 0, 'READY': 0}
    for it in items:
        st = str(it.get('status') or 'READY').upper()
        if 'UPLOADED' in st and 'NOT' not in st and 'PENDING' not in st:
            status_counts['UPLOADED'] += 1
        elif 'SUBMITTED' in st:
            status_counts['SUBMITTED'] += 1
        elif 'PENDING' in st:
            status_counts['PENDING UPLOAD'] += 1
        elif 'NOT' in st:
            status_counts['NOT UPLOADED'] += 1
        else:
            status_counts['READY'] += 1

    uploaded_pct = round((status_counts['UPLOADED'] / max(total_clips, 1)) * 100, 1)

    # ── SHEET 1: 📊 Summary & KPIs ──
    ws_sum = wb.active
    ws_sum.title = "📊 Summary & KPIs"
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_sum.merge_cells("A1:G1")
    c1 = ws_sum.cell(1, 1, "ONE TAP WORKSTATION — SAFAR MEDIA AUDIT & RECONCILIATION")
    c1.fill = brand_navy_fill
    c1.font = font_banner
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 36

    ws_sum.merge_cells("A2:G2")
    c2 = ws_sum.cell(2, 1, f"Archive: {title}  •  {subtitle}  •  Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    c2.fill = brand_blue_fill
    c2.font = font_subbanner
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 24

    # Top KPI Cards (Row 4 & 5)
    kpis = [
        ("A", "B", "TOTAL CLIPS", f"{total_clips:,}"),
        ("C", "D", "TOTAL RUNTIME", total_duration_fmt),
        ("E", "E", "STORAGE VOLUME", f"{total_size_gb} GB"),
        ("F", "G", "UPLOAD PROGRESS", f"{uploaded_pct}% ({status_counts['UPLOADED']}/{total_clips})")
    ]
    for start_c, end_c, label, val in kpis:
        col_s = openpyxl.utils.column_index_from_string(start_c)
        col_e = openpyxl.utils.column_index_from_string(end_c)
        if col_s != col_e:
            ws_sum.merge_cells(start_row=4, start_column=col_s, end_row=4, end_column=col_e)
            ws_sum.merge_cells(start_row=5, start_column=col_s, end_row=5, end_column=col_e)
        cl = ws_sum.cell(4, col_s, label)
        cl.font = font_kpi_label
        cl.alignment = Alignment(horizontal="center", vertical="center")
        cl.fill = kpi_card_fill
        cv = ws_sum.cell(5, col_s, val)
        cv.font = font_kpi_value
        cv.alignment = Alignment(horizontal="center", vertical="center")
        cv.fill = kpi_card_fill

    ws_sum.row_dimensions[4].height = 18
    ws_sum.row_dimensions[5].height = 28

    # Section 1: Event Segregation Breakdown (Row 7)
    ws_sum.merge_cells("A7:G7")
    cs1 = ws_sum.cell(7, 1, "1. EVENT SEGREGATION & KPI BREAKDOWN")
    cs1.font = font_section
    cs1.fill = section_fill
    cs1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_sum.row_dimensions[7].height = 26

    ev_headers = ["Event Code", "Official Event Name", "Total Clips", "Total Duration", "Data Size (MB)", "Uploaded Clips", "Upload %"]
    for col_idx, h in enumerate(ev_headers, 1):
        cell = ws_sum.cell(8, col_idx, h)
        cell.fill = dark_slate_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[8].height = 24

    # Group by Event Code
    events_grouped: Dict[str, List[Dict[str, Any]]] = {}
    for it in items:
        ev = (it.get('event') or '').strip().upper()
        events_grouped.setdefault(ev, []).append(it)

    sorted_ev_keys = sorted(events_grouped.keys(), key=lambda k: (-len(events_grouped[k]), k))

    current_r = 9
    for ev_code in sorted_ev_keys:
        ev_items = events_grouped[ev_code]
        ev_name = OFFICIAL_EVENTS_DICT.get(ev_code, ev_code if ev_code else "General / Uncategorized")
        ev_count = len(ev_items)
        ev_sec = sum(parse_duration_to_seconds(x.get('duration', '')) for x in ev_items)
        ev_dur = format_seconds_to_duration(ev_sec)
        ev_mb = round(sum((x.get('size_mb') or 0.0) for x in ev_items), 2)
        ev_up = sum(1 for x in ev_items if 'UPLOADED' in str(x.get('status') or '').upper() and 'NOT' not in str(x.get('status') or '').upper() and 'PENDING' not in str(x.get('status') or '').upper())
        ev_pct = f"{round((ev_up / max(ev_count, 1)) * 100, 1)}%"

        vals = [ev_code or '—', ev_name, ev_count, ev_dur, ev_mb, ev_up, ev_pct]
        is_even = (current_r % 2 == 0)
        for col_idx, val in enumerate(vals, 1):
            c = ws_sum.cell(current_r, col_idx, val)
            c.font = data_font
            c.border = thin_border
            if is_even:
                c.fill = alt_fill
            if col_idx in (1, 3, 4, 6, 7):
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        ws_sum.row_dimensions[current_r].height = 20
        current_r += 1

    # Total Row for Events
    tot_vals = ["TOTALS", f"{len(events_grouped)} Event Categories", total_clips, total_duration_fmt, round(total_size_mb, 2), status_counts['UPLOADED'], f"{uploaded_pct}%"]
    for col_idx, val in enumerate(tot_vals, 1):
        c = ws_sum.cell(current_r, col_idx, val)
        c.font = total_font
        c.border = total_border
        c.fill = kpi_card_fill
        if col_idx in (1, 3, 4, 6, 7):
            c.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == 5:
            c.alignment = Alignment(horizontal="right", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")
    ws_sum.row_dimensions[current_r].height = 22
    current_r += 2

    # Section 2: Chronological Date Timeline Breakdown
    ws_sum.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=7)
    cs2 = ws_sum.cell(current_r, 1, "2. CHRONOLOGICAL TIMELINE BREAKDOWN (DATE & SESSIONS)")
    cs2.font = font_section
    cs2.fill = section_fill
    cs2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_sum.row_dimensions[current_r].height = 26
    current_r += 1

    date_headers = ["Date", "Session / Subfolder", "Total Clips", "Total Duration", "Data Size (MB)", "Primary Events", "Status"]
    for col_idx, h in enumerate(date_headers, 1):
        cell = ws_sum.cell(current_r, col_idx, h)
        cell.fill = dark_slate_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[current_r].height = 24
    current_r += 1

    date_grouped: Dict[str, List[Dict[str, Any]]] = {}
    for it in items:
        d_key = it.get('date') or 'Undated'
        date_grouped.setdefault(d_key, []).append(it)

    for d_val in sorted(date_grouped.keys()):
        d_items = date_grouped[d_val]
        d_count = len(d_items)
        d_sec = sum(parse_duration_to_seconds(x.get('duration', '')) for x in d_items)
        d_dur = format_seconds_to_duration(d_sec)
        d_mb = round(sum((x.get('size_mb') or 0.0) for x in d_items), 2)
        d_evs = ", ".join(sorted(list(set((x.get('event') or '').strip() for x in d_items if (x.get('event') or '').strip()))))
        d_up = sum(1 for x in d_items if 'UPLOADED' in str(x.get('status') or '').upper() and 'NOT' not in str(x.get('status') or '').upper())
        d_st = f"{d_up}/{d_count} Uploaded"

        vals = [d_val, f"{d_count} files in date", d_count, d_dur, d_mb, d_evs or '—', d_st]
        is_even = (current_r % 2 == 0)
        for col_idx, val in enumerate(vals, 1):
            c = ws_sum.cell(current_r, col_idx, val)
            c.font = data_font
            c.border = thin_border
            if is_even:
                c.fill = alt_fill
            if col_idx in (1, 3, 4, 7):
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        ws_sum.row_dimensions[current_r].height = 20
        current_r += 1

    for col in ws_sum.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sum.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    # ── EVENT SEGREGATED SHEETS ──
    for ev_code in sorted_ev_keys:
        ev_items = events_grouped[ev_code]
        ev_name = OFFICIAL_EVENTS_DICT.get(ev_code, ev_code if ev_code else "Other")
        sheet_title = sanitize_sheet_title(f"{ev_code} - {ev_name}" if ev_code else "Other Events")
        
        ws_ev = wb.create_sheet(title=sheet_title)
        ws_ev.views.sheetView[0].showGridLines = True

        ws_ev.merge_cells("A1:H1")
        ch = ws_ev.cell(1, 1, f"EVENT CATEGORY: {ev_code} — {ev_name.upper()} ({len(ev_items)} CLIPS)")
        ch.fill = brand_navy_fill
        ch.font = font_banner
        ch.alignment = Alignment(horizontal="center", vertical="center")
        ws_ev.row_dimensions[1].height = 32

        headers = ["Date", "Session / Subfolder", "File Name", "Event", "Duration", "Size (MB)", "Status", "Remarks"]
        for col_idx, h in enumerate(headers, 1):
            c = ws_ev.cell(2, col_idx, h)
            c.fill = dark_slate_fill
            c.font = font_header
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws_ev.row_dimensions[2].height = 24

        ev_r = 3
        for item in sorted(ev_items, key=lambda x: (x.get('date', ''), x.get('file', ''))):
            raw_st = str(item.get('status') or 'READY').upper()
            if 'UPLOADED' in raw_st and 'NOT' not in raw_st and 'PENDING' not in raw_st:
                st_fill, st_font, st_lbl = fill_uploaded, font_uploaded, 'UPLOADED'
            elif 'SUBMITTED' in raw_st:
                st_fill, st_font, st_lbl = fill_submitted, font_submitted, 'SUBMITTED'
            elif 'PENDING' in raw_st:
                st_fill, st_font, st_lbl = fill_pending, font_pending, 'PENDING UPLOAD'
            elif 'NOT' in raw_st:
                st_fill, st_font, st_lbl = fill_not_uploaded, font_not_uploaded, 'NOT UPLOADED'
            else:
                st_fill, st_font, st_lbl = None, data_font, 'READY'

            row_vals = [
                item.get('date') or '',
                item.get('folder') or '',
                item.get('file') or '',
                item.get('event') or '',
                item.get('duration') or '',
                item.get('size_mb') if (item.get('size_mb') or 0) > 0 else '',
                st_lbl,
                item.get('remarks') or ''
            ]
            is_even = (ev_r % 2 == 0)
            for col_idx, val in enumerate(row_vals, 1):
                c = ws_ev.cell(ev_r, col_idx, val)
                c.font = data_font
                c.border = thin_border
                if is_even:
                    c.fill = alt_fill
                if col_idx == 7:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if st_fill:
                        c.fill = st_fill
                    if st_font:
                        c.font = st_font
                elif col_idx in (1, 4, 5):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 6:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
            ws_ev.row_dimensions[ev_r].height = 20
            ev_r += 1

        ev_sec = sum(parse_duration_to_seconds(x.get('duration', '')) for x in ev_items)
        ev_dur = format_seconds_to_duration(ev_sec)
        ev_mb = round(sum((x.get('size_mb') or 0.0) for x in ev_items), 2)
        sub_vals = ["SUBTOTAL", f"{len(ev_items)} files", "", ev_code, ev_dur, ev_mb, "", ""]
        for col_idx, val in enumerate(sub_vals, 1):
            c = ws_ev.cell(ev_r, col_idx, val)
            c.font = total_font
            c.border = total_border
            c.fill = kpi_card_fill
            if col_idx in (1, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        ws_ev.row_dimensions[ev_r].height = 22

        for col in ws_ev.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_ev.column_dimensions[col_letter].width = min(max(max_len + 3, 13), 50)

    # ── SHEET: Master Audit Sheet (All Items) ──
    ws_master = wb.create_sheet(title="Full Audit Sheet")
    ws_master.views.sheetView[0].showGridLines = True
    master_headers = ["Date", "Folder", "File", "Event", "Duration", "Size (MB)", "Status", "Remarks"]
    for col_idx, h in enumerate(master_headers, 1):
        cell = ws_master.cell(1, col_idx, h)
        cell.fill = brand_navy_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_master.row_dimensions[1].height = 28

    for r_idx, item in enumerate(items, 2):
        raw_st = str(item.get('status') or 'READY').upper()
        if 'UPLOADED' in raw_st and 'NOT' not in raw_st and 'PENDING' not in raw_st:
            st_fill, st_font, st_lbl = fill_uploaded, font_uploaded, 'UPLOADED'
        elif 'SUBMITTED' in raw_st:
            st_fill, st_font, st_lbl = fill_submitted, font_submitted, 'SUBMITTED'
        elif 'PENDING' in raw_st:
            st_fill, st_font, st_lbl = fill_pending, font_pending, 'PENDING UPLOAD'
        elif 'NOT' in raw_st:
            st_fill, st_font, st_lbl = fill_not_uploaded, font_not_uploaded, 'NOT UPLOADED'
        else:
            st_fill, st_font, st_lbl = None, data_font, 'READY'

        row_vals = [
            item.get('date') or '',
            item.get('folder') or '',
            item.get('file') or '',
            item.get('event') or '',
            item.get('duration') or '',
            item.get('size_mb') if (item.get('size_mb') or 0) > 0 else '',
            st_lbl,
            item.get('remarks') or ''
        ]
        is_even = (r_idx % 2 == 0)
        for col_idx, val in enumerate(row_vals, 1):
            c = ws_master.cell(r_idx, col_idx, val)
            c.font = data_font
            c.border = thin_border
            if is_even:
                c.fill = alt_fill
            if col_idx == 7:
                c.alignment = Alignment(horizontal="center", vertical="center")
                if st_fill:
                    c.fill = st_fill
                if st_font:
                    c.font = st_font
            elif col_idx in (1, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        ws_master.row_dimensions[r_idx].height = 20

    ws_master.auto_filter.ref = f"A1:H{len(items) + 1}"

    for col in ws_master.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_master.column_dimensions[col_letter].width = min(max(max_len + 4, 13), 55)

    # ── SHEET: Official Event Codes ──
    ws_codes = wb.create_sheet(title="Official Event Codes")
    ws_codes.views.sheetView[0].showGridLines = True
    ws_codes.row_dimensions[1].height = 26
    h2 = ["Abbreviation", "Official Event Name"]
    for col_idx, h in enumerate(h2, 1):
        c = ws_codes.cell(1, col_idx, h)
        c.fill = dark_slate_fill
        c.font = font_header
        c.alignment = Alignment(horizontal="center", vertical="center")

    for idx, (abbr, ev_name) in enumerate(OFFICIAL_EVENTS_LIST, 2):
        c1 = ws_codes.cell(idx, 1, abbr)
        c2 = ws_codes.cell(idx, 2, ev_name)
        c1.font = data_font
        c2.font = data_font
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = Alignment(horizontal="center", vertical="center")
        if idx % 2 == 0:
            c1.fill = alt_fill
            c2.fill = alt_fill
        ws_codes.row_dimensions[idx].height = 20
    ws_codes.column_dimensions['A'].width = 18
    ws_codes.column_dimensions['B'].width = 35

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)

    return {
        'success': True,
        'output_path': output_path,
        'total_rows': len(items),
        'safar_name': title,
        'total_duration': total_duration_fmt,
        'total_size_gb': total_size_gb,
        'events_count': len(events_grouped)
    }

def generate_query_audit_excel(
    query_text: str,
    items: List[Dict[str, Any]],
    output_path: str
) -> Dict[str, Any]:
    """
    Generates a dedicated query result Excel workbook for the Vault Assistant chat.
    """
    return build_segregated_workbook(
        title=f"Query: \"{query_text}\"",
        subtitle=f"Matched {len(items)} media clips across Archive Vault",
        items=items,
        output_path=output_path
    )

def generate_safar_audit_excel(
    archive_id: int,
    output_path: str,
    folder_path: Optional[str] = None,
    date_filter: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    db_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Generates a production-ready multi-tab segregated audit Excel workbook from SQLite.
    """
    conn = db_mod.get_connection(db_path)
    cur = conn.cursor()

    cur.execute("SELECT safar_name, city, safar_year FROM scanned_archives WHERE id = ?", (archive_id,))
    safar = cur.fetchone()
    safar_name = safar['safar_name'] if safar else f"Safar_{archive_id}"
    city = safar['city'] if safar else ""
    year = safar['safar_year'] if safar else ""

    query = """
        SELECT 
            m.file_name,
            m.media_type,
            m.size_bytes,
            m.duration_formatted,
            COALESCE(m.upload_status, 'unknown') as upload_status,
            f.relative_path,
            f.event_code,
            COALESCE(f.date_canonical_hijri, f.date_canonical_greg, '') as date_val
        FROM media_items m
        JOIN folder_nodes f ON f.id = m.folder_id
        WHERE m.archive_id = ?
    """
    params: List[Any] = [archive_id]
    if folder_path and folder_path.strip():
        fp = folder_path.strip()
        query += " AND (f.relative_path = ? OR f.relative_path LIKE ?)"
        params.extend([fp, f"{fp}%"])
    if date_filter and date_filter.strip():
        df = date_filter.strip()
        query += " AND (f.date_canonical_hijri = ? OR f.date_canonical_greg = ? OR f.relative_path LIKE ?)"
        params.extend([df, df, f"%{df}%"])
    elif date_from and date_to:
        query += " AND ((f.date_canonical_hijri BETWEEN ? AND ?) OR (f.date_canonical_greg BETWEEN ? AND ?))"
        params.extend([date_from, date_to, date_from, date_to])
    query += " ORDER BY f.relative_path, m.file_name"

    cur.execute(query, params)
    raw_rows = cur.fetchall()
    conn.close()

    items = []
    for r in raw_rows:
        size_mb = round((r['size_bytes'] or 0) / (1024 * 1024), 2)
        raw_st = (r['upload_status'] or 'unknown').lower()
        if raw_st == 'uploaded':
            st_label = 'UPLOADED'
        elif raw_st == 'submitted':
            st_label = 'SUBMITTED'
        elif raw_st == 'rendered_pending':
            st_label = 'PENDING UPLOAD'
        elif raw_st == 'not_uploaded':
            st_label = 'NOT UPLOADED'
        else:
            st_label = 'READY'

        items.append({
            'date': r['date_val'] or '',
            'folder': r['relative_path'] or '',
            'file': r['file_name'] or '',
            'event': r['event_code'] or '',
            'duration': r['duration_formatted'] or '',
            'size_mb': size_mb,
            'status': st_label,
            'remarks': ''
        })

    subtitle = f"{city} {year}".strip() or "Full Archive Audit"
    return build_segregated_workbook(
        title=safar_name,
        subtitle=subtitle,
        items=items,
        output_path=output_path
    )

def generate_raw_folder_audit_excel(
    folder_path: str,
    output_path: str,
    date_filter: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None
) -> Dict[str, Any]:
    """
    Crawls an unindexed external folder directly and produces a multi-tab segregated workbook.
    """
    if not os.path.exists(folder_path):
        raise FileNotFoundError(f"Folder not found on disk: {folder_path}")

    folder_path = os.path.normpath(folder_path)
    base_name = os.path.basename(folder_path) or 'Folder'

    items = []
    for root, dirs, files in os.walk(folder_path):
        dirs[:] = [
            d for d in dirs
            if not d.startswith('$')
            and d.lower() not in ('system volume information', '.trash', '.git', '.idea', '__pycache__', 'node_modules')
            and not scanner_mod.is_adobe_folder(d)
        ]

        folder_base = os.path.basename(root)
        if folder_base and scanner_mod.is_adobe_folder(folder_base):
            continue

        rel_dir = os.path.relpath(root, folder_path)
        if rel_dir == '.':
            rel_dir = ''

        dir_parts = rel_dir.split(os.sep) if rel_dir else []
        folder_date_val = ''
        folder_session = ''
        folder_event = ''

        for part in dir_parts:
            p_info = scanner_mod.parse_date_folder(part)
            if p_info.get('date_canonical_hijri'):
                folder_date_val = p_info['date_canonical_hijri']
            elif p_info.get('date_canonical_greg'):
                folder_date_val = p_info['date_canonical_greg']
            if p_info.get('session_tag'):
                folder_session = p_info['session_tag']
            ev = scanner_mod.extract_event_code(part)
            if ev:
                folder_event = ev

        full_date_label = f"{folder_date_val} {folder_session}".strip() if folder_date_val else ''

        if date_filter and date_filter.strip():
            df = date_filter.strip().lower()
            if df not in full_date_label.lower() and df not in rel_dir.lower():
                continue
        elif date_from and date_to and folder_date_val:
            if not (date_from <= folder_date_val <= date_to):
                continue

        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext not in scanner_mod.VIDEO_EXTS and ext not in scanner_mod.PHOTO_EXTS and ext not in scanner_mod.AUDIO_EXTS:
                continue
            if scanner_mod.is_adobe_file(f, ext):
                continue

            full_file_path = os.path.join(root, f)
            try:
                size_bytes = os.path.getsize(full_file_path)
            except Exception:
                size_bytes = 0
            size_mb = round(size_bytes / (1024 * 1024), 2)

            duration_formatted = ''
            if ext in scanner_mod.VIDEO_EXTS:
                try:
                    dur_sec = scanner_mod.get_mp4_duration_fast(full_file_path)
                    if dur_sec and dur_sec > 0:
                        duration_formatted = scanner_mod.format_duration(dur_sec)
                except Exception:
                    pass

            item_event = folder_event
            if not item_event:
                item_event = scanner_mod.extract_event_code(f)

            items.append({
                'date': folder_date_val,
                'folder': rel_dir,
                'file': f,
                'event': item_event,
                'duration': duration_formatted,
                'size_mb': size_mb,
                'status': 'READY',
                'remarks': ''
            })

    return build_segregated_workbook(
        title=base_name,
        subtitle="External Drive Folder Scan",
        items=items,
        output_path=output_path
    )


def generate_blank_audit_template(output_path: str) -> Dict[str, Any]:
    """Generates an official blank audit template Excel with instructions and event code list."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Sheet"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    sample_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = ["Date", "Folder", "File", "Event", "Status", "Remarks"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    sample_rows = [
        ["1448-03-13 _01M", "01-Mumbai-Bethak\\01-AQ\\Mulla Murtaza bhai Shaikh Juzer bhai", "C0001.MP4", "AQ", "Ready", "Sample entry"],
        ["1448-03-13 _01M", "01-Mumbai-Bethak\\01-AQ\\Mulla Murtaza bhai Shaikh Juzer bhai", "C0002.MP4", "AQ", "Ready", "Sample entry"],
        ["1448-03-13 _02E", "02-Mumbai-Majlis\\02-M\\Majlis Mubarak", "C0003.MP4", "M", "Ready", "Sample entry"]
    ]
    for r_idx, row in enumerate(sample_rows, 2):
        for col_idx, val in enumerate(row, 1):
            c = ws.cell(r_idx, col_idx, val)
            c.font = sample_font
            c.border = thin_border
            if col_idx in (1, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r_idx].height = 20

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 55)

    ws2 = wb.create_sheet(title="Official Event Codes")
    ws2.row_dimensions[1].height = 26
    h2 = ["Abbreviation", "Official Event Name"]
    for col_idx, h in enumerate(h2, 1):
        c = ws2.cell(1, col_idx, h)
        c.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for idx, (abbr, ev_name) in enumerate(OFFICIAL_EVENTS_LIST, 2):
        c1 = ws2.cell(idx, 1, abbr)
        c2 = ws2.cell(idx, 2, ev_name)
        c1.font = data_font
        c2.font = data_font
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = Alignment(horizontal="center", vertical="center")
        if idx % 2 == 0:
            c1.fill = alt_fill
            c2.fill = alt_fill
        ws2.row_dimensions[idx].height = 20
    ws2.column_dimensions['A'].width = 18
    ws2.column_dimensions['B'].width = 35

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)

    return {
        'success': True,
        'output_path': output_path
    }

def main():
    parser = argparse.ArgumentParser(description="One Tap Excel Reconciler CLI")
    parser.add_argument('--action', default='audit', choices=['audit', 'generate_safar', 'generate_external', 'generate_template', 'generate_query'], help="Action to perform")
    parser.add_argument('--excel', default=None, help="Path to Excel (.xlsx) file")
    parser.add_argument('--archive_id', type=int, default=None, help="Database Archive ID")
    parser.add_argument('--output', default=None, help="Output path for generated Excel")
    parser.add_argument('--sheet', default=None, help="Specific sheet name")
    parser.add_argument('--col_date', default=None, help="Name of date column")
    parser.add_argument('--col_folder', default=None, help="Name of folder column")
    parser.add_argument('--col_file', default=None, help="Name of filename column")
    parser.add_argument('--export', default=None, help="Path to export audit .xlsx")
    parser.add_argument('--folder_path', default=None, help="Specific folder path to generate Excel for")
    parser.add_argument('--folder', default=None, help="Alias for external folder path")
    parser.add_argument('--date_filter', default=None, help="Filter by date string or canonical session")
    parser.add_argument('--date_from', default=None, help="Filter start date")
    parser.add_argument('--date_to', default=None, help="Filter end date")
    parser.add_argument('--db', default=None, help="Custom SQLite DB path")
    parser.add_argument('--query', default=None, help="Search query string for query Excel generation")
    parser.add_argument('--items_json', default=None, help="JSON string or path to JSON file containing items to export")

    args = parser.parse_args()
    try:
        target_folder = args.folder_path or args.folder
        if args.action == 'generate_safar':
            if not args.archive_id or not args.output:
                print(json.dumps({'success': False, 'error': 'Missing --archive_id or --output'}))
                sys.exit(1)
            res = generate_safar_audit_excel(
                archive_id=args.archive_id,
                output_path=args.output,
                folder_path=target_folder,
                date_filter=args.date_filter,
                date_from=args.date_from,
                date_to=args.date_to,
                db_path=args.db
            )
            print(json.dumps(res))
        elif args.action == 'generate_external':
            if not target_folder or not args.output:
                print(json.dumps({'success': False, 'error': 'Missing --folder or --output for external generation'}))
                sys.exit(1)
            res = generate_raw_folder_audit_excel(
                folder_path=target_folder,
                output_path=args.output,
                date_filter=args.date_filter,
                date_from=args.date_from,
                date_to=args.date_to
            )
            print(json.dumps(res))
        elif args.action == 'generate_query':
            if not args.output:
                print(json.dumps({'success': False, 'error': 'Missing --output for generate_query'}))
                sys.exit(1)
            raw_items = []
            if args.items_json:
                if os.path.exists(args.items_json):
                    with open(args.items_json, 'r', encoding='utf-8') as jf:
                        raw_items = json.load(jf)
                else:
                    try:
                        raw_items = json.loads(args.items_json)
                    except Exception:
                        raw_items = []
            res = generate_query_audit_excel(
                query_text=args.query or "Custom Search Query",
                items=raw_items,
                output_path=args.output
            )
            print(json.dumps(res))
        elif args.action == 'generate_template':
            if not args.output:
                print(json.dumps({'success': False, 'error': 'Missing --output'}))
                sys.exit(1)
            res = generate_blank_audit_template(args.output)
            print(json.dumps(res))
        else:
            if not args.excel or not args.archive_id:
                print(json.dumps({'success': False, 'error': 'Missing --excel or --archive_id for audit'}))
                sys.exit(1)
            res = reconcile_excel_against_archive(
                excel_path=args.excel,
                archive_id=args.archive_id,
                sheet_name=args.sheet,
                col_date=args.col_date,
                col_folder=args.col_folder,
                col_file=args.col_file,
                date_filter=args.date_filter,
                date_from=args.date_from,
                date_to=args.date_to,
                db_path=args.db
            )

            if args.export:
                export_audit_excel(res, args.export)
                res['exported_file'] = args.export

            print(json.dumps(res))
    except Exception as e:
        print(json.dumps({'success': False, 'error': str(e)}))
        sys.exit(1)

if __name__ == '__main__':
    main()
